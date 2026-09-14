#!/usr/bin/env python3
"""
Automatic YCharts pull through the Excel Add-in (no manual import/export, no REST API).

Your YCharts plan includes the Excel Add-in, which pulls genuine YCharts numbers with the
=YCP(ticker, metric) and =YCS(ticker, metric, start) formulas. This tool builds a workbook full
of those formulas, then drives Excel (via xlwings) to refresh it and reads the results straight
into data/ycharts_cache/ + data/prices/ — the exact files the dashboard already uses. Run it on a
schedule and you have live YCharts data, hands-off.

REQUIREMENTS (all on your Mac, one-time):
  * Microsoft Excel (desktop).
  * The YCharts Excel Add-in installed and logged in (Account -> User ID + Access Key). You already
    use this — it's the same login you've been exporting from.
  * pip install xlwings openpyxl

USAGE:
  python -m ycharts_export.excel_bridge --watchlist sl_model --portfolio --build   # first time: make the workbook
  python -m ycharts_export.excel_bridge --refresh                                  # refresh + read into the cache
  # or both at once:
  python -m ycharts_export.excel_bridge --watchlist sl_model --portfolio --build --refresh

The workbook lives at data/ycharts_pull.xlsx. On a schedule (launchd/cron) --refresh gives you
live numbers. Deep history for charts can still come from Yahoo; the *current* YCharts values
(price, P/E, margins, targets...) and any YCS series come from YCharts here.
"""
from __future__ import annotations

import argparse
import os
import sys
import time
from datetime import date, datetime, timedelta

from .api import CACHE, PRICES, write, portfolio_tickers, SERIES_METRICS

WB = os.path.join("data", "ycharts_pull.xlsx")

# YCharts Excel metric codes for =YCP (current value). Edit freely; unknown codes just show an error cell.
POINT_METRICS = [
    "price", "market_cap", "pe_ratio", "forward_pe_ratio", "price_to_book_value",
    "price_to_sales_ratio", "ev_to_ebitda", "peg_ratio", "dividend_yield",
    "gross_profit_margin", "operating_margin", "profit_margin",
    "return_on_equity", "return_on_invested_capital", "free_cash_flow",
    "revenues", "net_income", "eps_diluted", "total_debt", "beta_5y",
    "analyst_target_price_mean", "analyst_recommendation",
]
# YCS series to pull as history (one sheet per ticker). Keep short — each is a spilled array.
SERIES_CODES = ["price", "pe_ratio", "revenues", "gross_profit_margin", "free_cash_flow"]


def _safe_sheet(sym):
    return ("S_" + sym)[:31].replace("/", "_")


def build_workbook(tickers, path=WB, years=15):
    """Create the formula workbook with openpyxl (no Excel needed to build)."""
    import openpyxl
    wb = openpyxl.Workbook()
    start = (date.today() - timedelta(days=int(years * 365.25))).isoformat()
    # points grid: A=ticker, row1 = metric codes, B2 = =YCP($A2, B$1)
    ws = wb.active; ws.title = "points"
    ws["A1"] = "Ticker"
    for j, m in enumerate(POINT_METRICS):
        ws.cell(row=1, column=2 + j, value=m)
    for i, t in enumerate(tickers):
        ws.cell(row=2 + i, column=1, value=t)
        for j, m in enumerate(POINT_METRICS):
            col = 2 + j
            ws.cell(row=2 + i, column=col,
                    value=f'=YCP($A{2 + i},{ws.cell(row=1, column=col).coordinate[:-1]}$1)')
    # one sheet per ticker: stepped YCS series (each spills date+value into 2 cols)
    for t in tickers:
        s = wb.create_sheet(_safe_sheet(t))
        s["A1"] = t
        for k, code in enumerate(SERIES_CODES):
            c = 1 + k * 3            # columns 1(A),4(D),7(G)... : date+value+gap
            s.cell(row=2, column=c, value=code)
            # YCS spills [date, value] starting the next row
            s.cell(row=3, column=c, value=f'=YCS("{t}","{code}","{start}")')
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print(f"built {path}: {len(tickers)} tickers, {len(POINT_METRICS)} point metrics, "
          f"{len(SERIES_CODES)} series each", file=sys.stderr)
    return path


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def refresh_and_read(path=WB, wait=180):
    """Open the workbook in Excel, let the YCharts add-in recalc, and read the values back."""
    try:
        import xlwings as xw
    except ImportError:
        sys.exit("xlwings is required for --refresh:  pip install xlwings")
    if not os.path.exists(path):
        sys.exit(f"{path} not found. Build it first: python -m ycharts_export.excel_bridge --build ...")
    app = xw.App(visible=False, add_book=False)
    try:
        wb = app.books.open(os.path.abspath(path), update_links=False)
        app.calculate()
        # The add-in fetches asynchronously; poll the points grid until it stops showing "loading"/errors.
        pts = wb.sheets["points"]
        deadline = time.time() + wait
        while time.time() < deadline:
            app.calculate()
            sample = pts.range("B2").value
            txt = str(sample).lower()
            if _num(sample) is not None or ("load" not in txt and "n/a" not in txt and sample not in (None, "")):
                time.sleep(3)  # let the rest settle
                break
            time.sleep(4)
        records = {}
        # points grid
        grid = pts.range("A1").expand().value
        if grid and isinstance(grid[0], list):
            metrics = grid[0][1:]
            for row in grid[1:]:
                sym = str(row[0]).upper().strip()
                if not sym:
                    continue
                rec = records.setdefault(sym, {"symbol": sym, "series": {}, "points": {},
                                               "as_of": datetime.now().isoformat(timespec="seconds"),
                                               "source": "ycharts (excel add-in)"})
                for m, val in zip(metrics, row[1:]):
                    n = _num(val)
                    if n is not None:
                        rec["points"][str(m)] = {"value": n}
        # per-ticker series sheets
        for sym in list(records):
            name = _safe_sheet(sym)
            if name not in [s.name for s in wb.sheets]:
                continue
            sh = wb.sheets[name]
            for k, code in enumerate(SERIES_CODES):
                col = 1 + k * 3
                block = sh.range((2, col)).expand("down").value  # header + spilled rows
                # block[0] is the code label; the YCS spill is a 2-col array we read directly
                arr = sh.range((3, col)).expand().value
                rows = []
                if isinstance(arr, list):
                    for r in arr:
                        if isinstance(r, (list, tuple)) and len(r) >= 2 and _num(r[1]) is not None:
                            d = r[0]
                            d = d.date().isoformat() if hasattr(d, "date") else str(d)[:10]
                            rows.append([d, _num(r[1])])
                if rows:
                    records[sym]["series"][code] = sorted(rows)
        wb.close()
        return records
    finally:
        app.quit()


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tickers", nargs="*")
    ap.add_argument("--watchlist")
    ap.add_argument("--portfolio", action="store_true")
    ap.add_argument("--build", action="store_true", help="(re)build the formula workbook")
    ap.add_argument("--refresh", action="store_true", help="open Excel, refresh the add-in, read into the cache")
    ap.add_argument("--years", type=int, default=15)
    ap.add_argument("--path", default=WB)
    args = ap.parse_args(argv)

    tickers = list(args.tickers)
    if args.watchlist:
        from portfolio.watchlists import load_user
        tickers += load_user(args.watchlist)
    if args.portfolio:
        tickers += portfolio_tickers()
    tickers = sorted(set(t.upper() for t in tickers))

    if args.build:
        if not tickers:
            sys.exit("give tickers / --watchlist / --portfolio to build the workbook")
        build_workbook(tickers, args.path, args.years)
    if args.refresh:
        print("Opening Excel and refreshing the YCharts add-in (keep Excel closed elsewhere)...", file=sys.stderr)
        records = refresh_and_read(args.path)
        write(records)
        n = sum(1 for r in records.values() if r.get("points") or r.get("series"))
        print(f"pulled live YCharts data for {n} tickers into {CACHE}/", file=sys.stderr)
    if not (args.build or args.refresh):
        ap.error("nothing to do: pass --build and/or --refresh")


if __name__ == "__main__":
    main()
