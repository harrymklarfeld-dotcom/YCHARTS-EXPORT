"""Unit-economics sanity tests for tenbagger/finance/model.py.

Run:  python -m pytest tenbagger/finance/tests -q
"""
import shutil
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

FIN = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(FIN))
import model as M  # noqa: E402

A = M.load_assumptions(FIN / "assumptions.yaml")
REVENUE_KEYS = ["rev_monthly", "rev_annual", "rev_student", "rev_pass", "rev_iap_gross", "rev_subs_net",
                "rev_ads", "rev_aff", "rev_b2b", "rev_total"]


@pytest.fixture(scope="module")
def res():
    return M.run_all(A)


def months(v, k):
    return v[k][1:]


# --- revenue and users --------------------------------------------------------
@pytest.mark.parametrize("scen", M.SCENARIOS)
def test_revenue_non_negative(res, scen):
    v = res[scen]["v"]
    for k in REVENUE_KEYS:
        assert min(months(v, k)) >= 0, k
    assert max(months(v, "refunds")) <= 0 and max(months(v, "store_fee")) <= 0


@pytest.mark.parametrize("scen", M.SCENARIOS)
def test_subs_never_exceed_mau(res, scen):
    v = res[scen]["v"]
    for m in range(1, M.MONTHS + 1):
        assert 0 <= v["subs_paid"][m] <= v["mau_total"][m]
        assert v["mau_free"][m] >= 0


def test_subs_never_exceed_mau_even_at_extreme_conversion():
    p = M.scenario_params(A, "bull", {"trial_start_rate": 1.0, "trial_to_paid": 1.0, "activation_rate": 1.0,
                                      "ret12_monthly": 1.0, "renew_annual": 1.0, "renew_student": 1.0})
    v = M.run_scenario(p)
    assert all(v["subs_paid"][m] <= v["mau_total"][m] + 1e-9 for m in range(1, 37))


def test_mix_over_100_percent_rejected():
    with pytest.raises(ValueError):
        M.scenario_params(A, "base", {"mix_monthly": 0.7, "mix_annual": 0.5})


# --- commission applied exactly once ---------------------------------------
def test_store_commission_applied_once():
    p = M.scenario_params(A, "base", {"refund_rate": 0.0, "store_fee": 0.15})
    v = M.run_scenario(p)
    for m in range(1, 37):
        gross = v["rev_iap_gross"][m]
        assert v["rev_subs_net"][m] == pytest.approx(gross * 0.85)
        assert v["store_fee"][m] == pytest.approx(-gross * 0.15)
        # store fee never touches ads, affiliate or B2B, and total = net subs + other lines
        assert v["rev_total"][m] == pytest.approx(v["rev_subs_net"][m] + v["rev_ads"][m] + v["rev_aff"][m]
                                                  + v["rev_b2b"][m])
    # commission is not also charged as a cost
    cogs_keys = [r.key for r in M.ROWS if r.key and r.key.startswith("cogs_") and r.key != "cogs_total"]
    assert "cogs_store" not in cogs_keys
    assert sum(v[k][36] for k in cogs_keys) == pytest.approx(v["cogs_total"][36])


def test_refund_then_commission_order():
    p = M.scenario_params(A, "base", {"refund_rate": 0.10, "store_fee": 0.15})
    v = M.run_scenario(p)
    g = v["rev_iap_gross"][24]
    assert v["rev_subs_net"][24] == pytest.approx(g * 0.9 * 0.85)
    assert p["net_factor"] == pytest.approx(0.9 * 0.85)


def test_revenuecat_fee_only_above_threshold():
    v = M.run_scenario(M.scenario_params(A, "base"))
    for m in range(1, 37):
        exp = max(0.0, v["bill_gross"][m] - 2500) * 0.01
        assert v["cogs_rc"][m] == pytest.approx(exp)


def test_ads_off_zeroes_ad_revenue_only():
    on = M.run_scenario(M.scenario_params(A, "base"))
    off = M.run_scenario(M.scenario_params(A, "base", {"ads_on": 0}))
    assert sum(months(off, "rev_ads")) == 0
    assert months(off, "rev_subs_net") == pytest.approx(months(on, "rev_subs_net"))
    assert off["ebit"][36] == pytest.approx(on["ebit"][36] - on["rev_ads"][36])


def test_free_users_capped_at_one_linked_bank():
    v = M.run_scenario(M.scenario_params(A, "base", {"plaid_free_link_share": 1.0}))
    assert all(v["items_free"][m] <= v["mau_free"][m] + 1e-9 for m in range(1, 37))
    assert v["items_free"][12] == 0  # linking launches in month 13 in the base case


# --- break-even detection -------------------------------------------------
def test_breakeven_detection_synthetic():
    series = [None, -5, -3, 1, -1, 2, 3, 4]
    assert M.breakeven(series) == {"first": 3, "sustained": 5}
    assert M.breakeven([None, -1, -2, -3]) == {"first": "Never", "sustained": "Never"}
    assert M.breakeven([None, 0, 1, 2]) == {"first": 1, "sustained": 1}
    assert M.first_month([None, 0, 0, 1, 1]) == 3
    assert M.first_month([None, 0, 0]) == "Never"


def test_breakeven_matches_flags(res):
    for s in M.SCENARIOS:
        v, k = res[s]["v"], res[s]["kpi"]
        be = M.breakeven(v["ebit"])
        assert k["First month with EBIT >= 0"] == be["first"]
        assert k["Break-even month (EBIT >= 0 from then on)"] == be["sustained"]


def test_breakeven_responds_to_price():
    rich = M.run_scenario(M.scenario_params(A, "base", {"price_monthly": 200, "price_annual": 2000,
                                                         "price_student": 1000}))
    zero = M.run_scenario(M.scenario_params(A, "base", {"trial_start_rate": 0, "ads_on": 0, "affiliate_on": 0}))
    assert M.breakeven(rich["ebit"])["sustained"] != "Never"
    assert M.breakeven(zero["ebit"])["first"] == "Never"


# --- yaml overrides -------------------------------------------------------
def test_scenario_overrides_applied():
    base = M.scenario_params(A, "base")
    bear = M.scenario_params(A, "bear")
    assert base["price_annual"] == 79.99 and bear["price_annual"] == 59.99
    assert bear["store_fee"] == base["store_fee"]  # not overridden -> follows base
    assert M.scenario_params(A, "base", {"asa_cpi": 99})["asa_cpi"] == 99


def test_unknown_override_rejected(tmp_path):
    raw = yaml.safe_load(open(FIN / "assumptions.yaml"))
    raw["scenarios"]["bear"]["not_a_driver"] = 1
    f = tmp_path / "a.yaml"
    f.write_text(yaml.safe_dump(raw))
    with pytest.raises(KeyError):
        M.load_assumptions(f)
    with pytest.raises(KeyError):
        M.scenario_params(A, "base", {"typo_key": 1})


def test_yaml_edit_changes_results(tmp_path):
    raw = yaml.safe_load(open(FIN / "assumptions.yaml"))
    raw["drivers"]["Pricing"]["price_annual"]["v"] = 99.99
    f = tmp_path / "a.yaml"
    f.write_text(yaml.safe_dump(raw))
    A2 = M.load_assumptions(f)
    r1 = M.run_scenario(M.scenario_params(A, "base"))
    r2 = M.run_scenario(M.scenario_params(A2, "base"))
    assert r2["rev_annual"][36] == pytest.approx(r1["rev_annual"][36] * 99.99 / 79.99)


def test_every_driver_has_source():
    for d in A["drivers"].values():
        assert d.src.strip() and d.label.strip()


# --- workbook ---------------------------------------------------------------
def test_workbook_has_live_formulas(tmp_path, res):
    sens = M.sensitivities(A)
    path = tmp_path / "m.xlsx"
    expect = M.build_workbook(A, res, sens, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Assumptions", "Monthly_Base", "Monthly_Bear", "Monthly_Bull", "Annual_PnL", "KPIs",
                             "Sensitivity", "Sources"]
    ws = wb["Monthly_Base"]
    assert str(ws["C6"].value).startswith("=") and "Assumptions!" in ws["C6"].value
    a = wb["Assumptions"]
    assert a["E13"].font.color.rgb.endswith("0000FF")      # inputs blue
    assert ws["C6"].font.color is None or ws["C6"].font.color.rgb.endswith("000000")  # formulas black
    assert ws.freeze_panes == "C5"
    assert len(expect) > 5000


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_libreoffice_recalc_matches_python(tmp_path, res):
    sens = M.sensitivities(A)
    path = tmp_path / "m.xlsx"
    expect = M.build_workbook(A, res, sens, path)
    ver = M.verify_with_libreoffice(path, expect)
    if ver["status"] == "skipped":
        pytest.skip(ver["reason"])
    assert ver["status"] == "ok", ver["examples"]
    assert ver["formula_errors"] == 0
